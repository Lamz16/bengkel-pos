"""Editable Word export for the generated client guide."""
from __future__ import annotations
import re
from pathlib import Path

def export_docx(output: Path) -> Path:
    try:
        from docx import Document
        from docx.enum.style import WD_STYLE_TYPE
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt, Cm
    except ImportError as exc:
        raise RuntimeError("Ekspor DOCX memerlukan python-docx dari runtime workspace.") from exc
    source=output / "manual-guide-draft.md"
    if not source.exists(): raise FileNotFoundError("Jalankan generate terlebih dahulu.")
    doc=Document(); section=doc.sections[0]
    section.top_margin=section.bottom_margin=Cm(2); section.left_margin=section.right_margin=Cm(2.2)
    normal=doc.styles["Normal"]; normal.font.name="Aptos"; normal.font.size=Pt(10.5)
    for name,size in (("Title",24),("Heading 1",16),("Heading 2",13),("Heading 3",11)):
        style=doc.styles[name]; style.font.name="Aptos"; style.font.size=Pt(size); style.font.color.rgb=None
    if "Screenshot Placeholder" not in [s.name for s in doc.styles]:
        style=doc.styles.add_style("Screenshot Placeholder", WD_STYLE_TYPE.PARAGRAPH); style.font.name="Aptos"; style.font.size=Pt(10); style.font.italic=True
    for line in source.read_text(encoding="utf-8").splitlines():
        if not line or line.startswith("<!--"): continue
        if line.startswith("# "):
            p=doc.add_paragraph(line[2:], style="Title"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith("## "): doc.add_paragraph(line[3:], style="Heading 1")
        elif line.startswith("### "): doc.add_paragraph(line[4:], style="Heading 2")
        elif line.startswith("- "): doc.add_paragraph(line[2:], style="List Bullet")
        elif re.match(r"\d+\. ", line): doc.add_paragraph(re.sub(r"^\d+\. ", "", line), style="List Number")
        elif line.startswith("> 📷 Screenshot"): doc.add_paragraph(line[2:], style="Screenshot Placeholder")
        elif line == "---": doc.add_paragraph()
        else:
            # Make common Markdown emphasis readable but keep the text editable.
            p=doc.add_paragraph(); chunks=re.split(r"(\*\*[^*]+\*\*)", line)
            for chunk in chunks:
                run=p.add_run(chunk.strip("*")); run.bold=chunk.startswith("**")
    target=output / "Panduan Penggunaan BengkelPro POS.docx"; doc.save(target); return target
